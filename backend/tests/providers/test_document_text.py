import asyncio
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from interview_api.domain.errors import InvalidDocumentError, UnsupportedDocumentTypeError
from interview_api.domain.models import DocumentInput
from interview_api.providers.document_text import PdfDocxTextExtractor


def make_document(content: bytes, filename: str) -> DocumentInput:
    return DocumentInput(
        stream=BytesIO(content),
        size_bytes=len(content),
        context_type="cv",
        file_id="cv",
        filename=filename,
    )


def test_extracts_utf8_text() -> None:
    extractor = PdfDocxTextExtractor()
    text = asyncio.run(extractor.extract(make_document("Опыт Python".encode(), "cv.txt")))
    assert text == "Опыт Python"


def test_extracts_docx_paragraphs_and_tables() -> None:
    stream = BytesIO()
    document = DocxDocument()
    document.add_paragraph("Python developer")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "FastAPI"
    document.save(stream)
    extractor = PdfDocxTextExtractor()

    text = asyncio.run(extractor.extract(make_document(stream.getvalue(), "cv.docx")))

    assert text == "Python developer\nFastAPI"


def test_extracts_pdf_text() -> None:
    stream = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 200 Td (Python developer) Tj ET")
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(stream)
    extractor = PdfDocxTextExtractor()

    text = asyncio.run(extractor.extract(make_document(stream.getvalue(), "vacancy.pdf")))

    assert text == "Python developer"


def test_rejects_pdf_without_text(monkeypatch) -> None:
    monkeypatch.setattr(
        "interview_api.providers.document_text.PdfReader",
        lambda _stream: SimpleNamespace(
            is_encrypted=False,
            pages=[SimpleNamespace(extract_text=lambda: "")],
        ),
    )
    extractor = PdfDocxTextExtractor()

    with pytest.raises(InvalidDocumentError):
        asyncio.run(extractor.extract(make_document(b"pdf", "scan.pdf")))


def test_rejects_unsupported_extension() -> None:
    extractor = PdfDocxTextExtractor()

    with pytest.raises(UnsupportedDocumentTypeError):
        asyncio.run(extractor.extract(make_document(b"text", "cv.rtf")))

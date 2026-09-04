import asyncio
from pathlib import Path
from zipfile import BadZipFile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from interview_api.domain.errors import InvalidDocumentError, UnsupportedDocumentTypeError
from interview_api.domain.models import DocumentInput


class PdfDocxTextExtractor:
    """Extract plain text from a request-scoped PDF, DOCX, or TXT stream."""

    async def extract(self, document: DocumentInput) -> str:
        try:
            text = await asyncio.to_thread(self._extract_sync, document)
        except UnsupportedDocumentTypeError:
            raise
        except (
            BadZipFile,
            OSError,
            PackageNotFoundError,
            PyPdfError,
            UnicodeError,
            ValueError,
        ) as exc:
            raise InvalidDocumentError(document.filename) from exc

        normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        if not normalized:
            raise InvalidDocumentError(document.filename)
        return normalized

    def _extract_sync(self, document: DocumentInput) -> str:
        suffix = Path(document.filename or "").suffix.casefold()
        document.stream.seek(0)

        if suffix == ".pdf":
            reader = PdfReader(document.stream)
            if reader.is_encrypted:
                raise InvalidDocumentError(document.filename)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix == ".docx":
            docx = DocxDocument(document.stream)
            paragraphs = [paragraph.text for paragraph in docx.paragraphs]
            table_cells = [
                cell.text
                for table in docx.tables
                for row in table.rows
                for cell in row.cells
            ]
            return "\n".join([*paragraphs, *table_cells])
        if suffix == ".txt":
            return document.stream.read().decode("utf-8-sig")

        raise UnsupportedDocumentTypeError(document.filename)
